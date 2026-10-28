"""Build kb/fitline-places.xlsx from kb/data/*.csv (run build_places.py first)."""
import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "data"
OUT = ROOT / "fitline-places.xlsx"

INK = "10191B"
CEDAR = "0E4F47"
HEAD_FILL = PatternFill("solid", fgColor=CEDAR)
HEAD_FONT = Font(bold=True, color="FFFFFF")
PU_FILL = PatternFill("solid", fgColor="E3F1EC")
CHECK_FILL = PatternFill("solid", fgColor="FFF1D6")


def read(name):
    with open(DATA / name, encoding="utf-8") as f:
        return list(csv.DictReader(f))


# ---- Facts kept by hand: corrections and country-level travel facts Wikidata does not give cleanly ----
COUNTRY_EXTRA = {
    "FR": dict(type="Sovereign state", un_region="Western Europe", languages="fr", timezones="Europe/Paris", plugs="C, E", levels="Région › [Area] › Planning unit", notes="Metropolitan France only for now; overseas regions later. Monaco is its own country."),
    "ES": dict(type="Sovereign state", un_region="Southern Europe", languages="es (+ ca, eu, gl regionally)", timezones="Europe/Madrid; Atlantic/Canary", plugs="C, F", levels="Autonomous community › [Coast] › Planning unit", notes=""),
    "US": dict(type="Sovereign state", un_region="Northern America", languages="en", timezones="America/New_York; America/Chicago; America/Denver; America/Phoenix; America/Los_Angeles; America/Anchorage; Pacific/Honolulu", plugs="A, B", levels="State › [Travel region] › Planning unit", notes="Territories (Puerto Rico, Guam...) later, as their own entries."),
    "IT": dict(type="Sovereign state", un_region="Southern Europe", languages="it", timezones="Europe/Rome", plugs="C, F, L", levels="Region › [Area] › Planning unit", notes="Vatican City and San Marino are their own countries."),
    "CN": dict(type="Sovereign state", un_region="Eastern Asia", languages="zh", timezones="Asia/Shanghai", plugs="A, C, I", levels="Province-level division › Planning unit", notes="Hong Kong, Macau and Taiwan are separate ISO entries, added later."),
    "TR": dict(type="Sovereign state", un_region="Western Asia", languages="tr", timezones="Europe/Istanbul", plugs="C, F", levels="Geographic region › Planning unit", notes="Official English name Türkiye (ISO); travellers still search 'Turkey'."),
    "MX": dict(type="Sovereign state", un_region="Central America", languages="es", timezones="America/Mexico_City; America/Cancun; America/Merida; America/Mazatlan; America/Chihuahua; America/Tijuana", plugs="A, B", levels="State › [Coast] › Planning unit", notes=""),
    "GB": dict(type="Sovereign state", un_region="Northern Europe", languages="en (+ cy, gd regionally)", timezones="Europe/London", plugs="G", levels="Nation › [Region, England only] › Planning unit", notes="England's regions are the official statistical regions, used as travel regions."),
    "DE": dict(type="Sovereign state", un_region="Western Europe", languages="de", timezones="Europe/Berlin", plugs="C, F", levels="State › Planning unit", notes="City-states Berlin, Hamburg and Bremen are their own planning unit."),
    "GR": dict(type="Sovereign state", un_region="Southern Europe", languages="el", timezones="Europe/Athens", plugs="C, F", levels="Region › [Island group] › Planning unit", notes=""),
    "JP": dict(type="Sovereign state", un_region="Eastern Asia", languages="ja", timezones="Asia/Tokyo", plugs="A, B", levels="Region › Prefecture › Planning unit", notes="The 8 regions are traditional, not official."),
    "TH": dict(type="Sovereign state", un_region="South-eastern Asia", languages="th", timezones="Asia/Bangkok", plugs="A, B, C, O", levels="Region › Planning unit", notes="Regions are travel regions; the 77 provinces are not a level."),
    "SG": dict(type="Sovereign state", un_region="South-eastern Asia", languages="en, ms, zh, ta", timezones="Asia/Singapore", plugs="G", levels="Country is the planning unit", notes=""),
}

LEVEL_TYPES = [
    ("FR", 1, "Région", "Région", "admin", ""), ("FR", 2, "Area", "", "travel", "Optional: Côte d'Azur, Provence"),
    ("ES", 1, "Autonomous community", "Comunidad autónoma", "admin", ""), ("ES", 2, "Coast", "Costa", "travel", "Optional: Costa del Sol"),
    ("US", 1, "State", "", "admin", "Plus the federal district (Washington, D.C.)"), ("US", 2, "Travel region", "", "travel", "Optional: California's regions"),
    ("IT", 1, "Region", "Regione", "admin", ""), ("IT", 2, "Area", "", "travel", "Optional: Dolomites"),
    ("CN", 1, "Province-level division", "省级行政区", "admin", "Provinces, municipalities, autonomous regions"),
    ("TR", 1, "Geographic region", "Coğrafi bölge", "travel", "The 7 standard regions"),
    ("MX", 1, "State", "Estado", "admin", "Plus Mexico City"), ("MX", 2, "Coast", "", "travel", "Optional: Riviera Maya"),
    ("GB", 1, "Nation", "", "admin", "England, Scotland, Wales, Northern Ireland"), ("GB", 2, "Region", "", "travel", "England only: official statistical regions"),
    ("DE", 1, "State", "Land", "admin", ""),
    ("GR", 1, "Region", "Periféreia", "admin", ""), ("GR", 2, "Island group", "", "travel", "Optional: Cyclades, Dodecanese"),
    ("JP", 1, "Region", "地方", "travel", "The 8 traditional regions"), ("JP", 2, "Prefecture", "都道府県", "admin", ""),
    ("TH", 1, "Region", "ภาค", "travel", "Travel regions; provinces are not a level"),
]

US_TZ = {**{s: "America/New_York" for s in "CT DE DC FL GA ME MD MA NH NJ NY NC OH PA RI SC VT VA WV KY".split()},
         **{s: "America/Chicago" for s in "AL AR IL IA LA MN MS MO OK WI TX TN KS NE ND SD".split()},
         **{s: "America/Denver" for s in "CO MT NM UT WY ID".split()},
         "AZ": "America/Phoenix", "MI": "America/Detroit", "IN": "America/Indiana/Indianapolis",
         **{s: "America/Los_Angeles" for s in "CA NV OR WA".split()}, "AK": "America/Anchorage", "HI": "Pacific/Honolulu"}
MX_TZ = {"ROO": "America/Cancun", "YUC": "America/Merida", "CAM": "America/Merida", "BCS": "America/Mazatlan", "SIN": "America/Mazatlan",
         "NAY": "America/Mazatlan", "CHH": "America/Chihuahua", "BCN": "America/Tijuana", "SON": "America/Hermosillo"}
TZ_OVERRIDE = {"US-P-GREAT-SMOKY-MOUNTAINS": "America/New_York", "US-P-BLACK-HILLS": "America/Denver", "MX-P-SAYULITA-RIVIERA-NAYARIT": "America/Mexico_City"}
SINGLE_TZ = {k: v["timezones"] for k, v in COUNTRY_EXTRA.items() if ";" not in v["timezones"]}


def timezone_of(r, rows):
    if r["place_id"] in TZ_OVERRIDE:
        return TZ_OVERRIDE[r["place_id"]]
    iso = r["country_iso2"]
    if iso in SINGLE_TZ:
        return SINGLE_TZ[iso]
    chain, p = [], r
    while p:
        chain.append(p["place_id"])
        p = rows.get(p["parent_id"])
    if iso == "ES":
        return "Atlantic/Canary" if "ES-CN" in chain else "Europe/Madrid"
    sub = next((c for c in chain if c.startswith(iso + "-") and len(c.split("-")) == 2), None)
    if not sub:
        return ""
    code = sub.split("-")[1]
    return US_TZ.get(code, "") if iso == "US" else MX_TZ.get(code, "America/Mexico_City")


def style_sheet(ws, widths, freeze="B2"):
    for c in ws[1]:
        c.fill, c.font = HEAD_FILL, HEAD_FONT
        c.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 22
    ws.freeze_panes = freeze
    ws.auto_filter.ref = ws.dimensions
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def main():
    places = read("places.csv")
    rows = {r["place_id"]: r for r in places}
    countries = read("countries.csv")
    links = read("place_links.csv")
    report = {r["place_id"]: r for r in read("match_report.csv")}

    for r in places:
        r["timezone"] = timezone_of(r, rows)
        check = report.get(r["place_id"], {}).get("why", "")
        if r["source"] == "Curated":
            r["status"] = "draft" if check == "country + coordinates" else "needs check"
            r["match_note"] = "" if check == "country + coordinates" else (check or "no Wikidata match")
        else:
            r["match_note"] = ""

    wb = Workbook()

    # README
    ws = wb.active
    ws.title = "README"
    readme = [
        ("Fitline places: knowledge base, level 1", ""),
        ("", ""),
        ("What this is", "Every place a traveller might pick, from country down to the 'planning unit': the place you would base yourself, where attractions are grouped as the same place and shown as pins on the map."),
        ("Scope (v1)", "Tier 1: France, Spain, United States, Italy, China, Türkiye, Mexico, United Kingdom, Germany, Greece, Japan, Thailand, Singapore."),
        ("How the tree works", "One row per place, linked to its parent by parent_id. Depth varies by country: Singapore has 1 level, Japan 3. The hierarchy follows how travellers think: official regions where they match, travel regions (Kantō, the Amalfi Coast, the Cyclades) where they don't."),
        ("", ""),
        ("Key columns", ""),
        ("place_id", "Stable ID, never reused. Countries use ISO 3166-1 (JP). Official regions use ISO 3166-2 (JP-13). Travel regions are XX-R-NAME, planning units XX-P-NAME."),
        ("parent_id", "The one main parent, used to drill down on the map. Places that also belong elsewhere are listed in place_links."),
        ("kind", "admin = an official unit (its outline comes from official boundaries). travel = a travel region or planning unit (outline merged from official ones or drawn later)."),
        ("is_planning_unit", "yes = attractions show as pins at this level. Everything above it shows count bubbles instead."),
        ("day_trip_from", "For places usually visited from another base (Versailles from Paris), the place_id of that base."),
        ("wikidata / osm_relation", "Cross-references for names in every language, photos and map outlines."),
        ("default_zoom", "Starting map zoom for this place; a first guess by kind of place, adjust per row."),
        ("status", "reviewed = official data from Wikidata. draft = curated and matched to Wikidata. needs check = the Wikidata match is doubtful or missing; see match_note."),
        ("", ""),
        ("Sources", "Official names, codes, coordinates, local-script names and map IDs: Wikidata (CC0). Travel regions and planning units: curated for this project."),
        ("How to change it", "Edit kb/data/curated/<country>.txt and rerun kb/scripts/build_places.py then build_workbook.py, or edit rows here and copy them back."),
        ("Not in v1", "French overseas regions; Hong Kong, Macau, Taiwan; Monaco, Vatican City, San Marino; US territories; neighbourhoods below planning units; map outlines for travel regions."),
    ]
    for a, b in readme:
        ws.append([a, b])
    ws["A1"].font = Font(bold=True, size=14, color=CEDAR)
    for row in ws.iter_rows(min_row=3):
        row[0].font = Font(bold=True, color=INK)
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
        row[0].alignment = Alignment(vertical="top")
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 110

    # countries
    ws = wb.create_sheet("countries")
    head = ["iso2", "iso3", "name_en", "name_local", "type", "un_region", "capital", "currency", "languages", "timezones",
            "driving_side", "plug_types", "levels_used", "planning_units", "wikidata", "osm_relation", "lat", "lon", "notes"]
    ws.append(head)
    pu_count = {}
    for r in places:
        if r["is_planning_unit"]:
            pu_count[r["country_iso2"]] = pu_count.get(r["country_iso2"], 0) + 1
    names = {r["place_id"]: r["name_en"] for r in places if r["level_type"] == "Country"}
    for c in countries:
        x = COUNTRY_EXTRA[c["iso2"]]
        ws.append([c["iso2"], c["iso3"], names[c["iso2"]], c["name_local"], x["type"], x["un_region"], c["capital"],
                   c["currency"].split(";")[0], x["languages"], x["timezones"], c["driving_side"], x["plugs"], x["levels"],
                   pu_count.get(c["iso2"], 0), c["wikidata"], c["osm_relation"], float(c["lat"]), float(c["lon"]), x["notes"]])
    style_sheet(ws, [6, 6, 16, 18, 15, 18, 16, 9, 22, 40, 11, 11, 42, 13, 10, 12, 10, 10, 60], freeze="D2")

    # level_types
    ws = wb.create_sheet("level_types")
    ws.append(["iso2", "depth", "level_name_en", "level_name_local", "kind", "notes"])
    for t in LEVEL_TYPES:
        ws.append(list(t))
    style_sheet(ws, [6, 7, 24, 20, 8, 46])

    # places
    ws = wb.create_sheet("places")
    cols = ["place_id", "parent_id", "country_iso2", "depth", "level_type", "kind", "name_en", "name_local", "path",
            "is_planning_unit", "day_trip_from", "iso_3166_2", "wikidata", "osm_relation", "lat", "lon", "default_zoom",
            "timezone", "slug", "notes", "source", "status", "match_note"]
    ws.append(cols)
    for r in places:
        out = []
        for col in cols:
            v = r.get(col, "")
            if col in ("depth", "default_zoom") and v != "":
                v = int(v)
            elif col in ("lat", "lon") and v != "":
                v = float(v)
            out.append(v)
        ws.append(out)
    style_sheet(ws, [26, 22, 8, 6, 18, 7, 30, 22, 58, 9, 26, 10, 11, 11, 10, 10, 7, 20, 26, 50, 9, 11, 28], freeze="C2")
    n = len(places) + 1
    pu_col = get_column_letter(cols.index("is_planning_unit") + 1)
    st_col = get_column_letter(cols.index("status") + 1)
    ws.conditional_formatting.add(f"A2:{get_column_letter(len(cols))}{n}", FormulaRule(formula=[f'${st_col}2="needs check"'], fill=CHECK_FILL))
    ws.conditional_formatting.add(f"A2:{get_column_letter(len(cols))}{n}", FormulaRule(formula=[f'${pu_col}2="yes"'], fill=PU_FILL))

    # place_links
    ws = wb.create_sheet("place_links")
    ws.append(["place_id", "place_name", "linked_place_id", "linked_name", "relation"])
    for l in links:
        ws.append([l["place_id"], rows[l["place_id"]]["name_en"], l["linked_place_id"], rows[l["linked_place_id"]]["name_en"], l["relation"]])
    style_sheet(ws, [30, 32, 24, 28, 12])

    # vocab (dropdown sources)
    ws = wb.create_sheet("vocab")
    level_types = sorted({r["level_type"] for r in places})
    vocab = {
        "kind": ["admin", "travel"],
        "status": ["reviewed", "draft", "needs check"],
        "is_planning_unit": ["yes", ""],
        "level_type": level_types,
        "country_iso2": [c["iso2"] for c in countries],
    }
    for i, (k, values) in enumerate(vocab.items(), 1):
        ws.cell(row=1, column=i, value=k)
        for j, v in enumerate(values, 2):
            ws.cell(row=j, column=i, value=v)
    style_sheet(ws, [12, 14, 16, 24, 14], freeze="A2")

    wsp = wb["places"]
    for k, col_name in [("kind", "kind"), ("status", "status"), ("is_planning_unit", "is_planning_unit"), ("level_type", "level_type"), ("country_iso2", "country_iso2")]:
        vi = list(vocab).index(k) + 1
        letter = get_column_letter(vi)
        dv = DataValidation(type="list", formula1=f"=vocab!${letter}$2:${letter}${len(vocab[k]) + 1}", allow_blank=True)
        col = get_column_letter(cols.index(col_name) + 1)
        dv.add(f"{col}2:{col}{n + 500}")
        wsp.add_data_validation(dv)

    wb.save(OUT)
    print("saved", OUT, "rows", len(places))


if __name__ == "__main__":
    main()
